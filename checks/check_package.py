"""High-signal consistency checks for the  submission package."""
import re
import sys
import zipfile
import hashlib
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parent
MS = ROOT / "manuscript/SPOT_giant_virus_Manuscript-.docx"
CL = ROOT / "manuscript/Cover_Letter-Nature-.docx"
SI = ROOT / "supplementary/Supplementary_Information.docx"
GD = ROOT / "supplementary/SIGuide.docx"
SF = ROOT / "supplementary/Supplementary_Figures.docx"
checks = []


def doc_text(path):
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] +
                      [c.text for t in doc.tables for r in t.rows for c in r.cells])


def check(label, condition, detail=""):
    checks.append(bool(condition))
    print(f"[{'PASS' if condition else 'FAIL'}] {label}" + (f" — {detail}" if detail else ""))


files = [MS, CL, SI, GD, SF]
check("five  documents exist", all(p.exists() for p in files))
ms, cl, si, gd, sf = map(doc_text, files)

tables = sorted(ROOT.glob("supplementary/Supplementary_Table_*.xlsx"))
numbers = sorted({int(re.search(r"Table_(\d+)_", p.name).group(1)) for p in tables})
expected_tables = list(range(1, 11)) + list(range(12, 19))
check("all analysis-generated tables are present (Table 11 remains author-completed)",
      numbers == expected_tables, str(numbers))

summary = ms[ms.index("Ocean warming is reorganizing"):ms.index("Marine species are shifting")]
summary_words = len(re.sub(r"\d", "", summary).split())
check("Summary is at most 200 words", summary_words <= 200, str(summary_words))
check("main-text references remain within Nature's 50-reference limit",
      "39 main-text references" not in ms and len(re.findall(r"^\d+\. [A-Z]", ms[:ms.index("Figure legends")], re.M)) == 39)

tara = pd.read_csv(ROOT.parent / "data/Tara_observed_Dindex.csv")
tara = tara[tara.D_index_qc.eq("ok")]
q = tara.I_plus_A.quantile([0.25, 0.5, 0.75]) * 100
check("Tara dominance statistic reproduces from 196 valid samples",
      len(tara) == 196 and round(q.loc[0.5], 1) == 96.9 and round(q.loc[0.25], 1) == 91.7 and round(q.loc[0.75], 1) == 99.0)
check("manuscript reports the reproduced Tara statistic",
      "median 96.9%" in ms and "91.7–99.0%" in ms)

check("overclaim and deprecated term removed",
      "independently of genome recovery" not in ms and "thermophilization" not in ms)
check("unweighted support diagnostic is not called canonical AOA",
      "not presented as the canonical AOA method" in ms and "area-of-applicability results" not in ms)
#  replaced the  disclosure with the tested pre-2014 derivation plus the residual caveat
check("discovery/test dependence is tested and the residual caveat disclosed",
      "restricted to the 89 metagenomes collected before 2014" in ms and "made with knowledge of the full series" in ms)
check("SSP5-8.5 is identified as a high-emissions stress test",
      "high-emissions stress test" in ms and "high-emissions stress test" in cl)
check("equal-dominance front is defined in the Fig. 5 legend",
      "D = 0 equal-dominance fronts" in ms)

check("current Supplementary Fig. 3 is the projected-components figure",
      "Supplementary Fig. 3 | Projected Imitervirales" in si and
      "Supplementary Fig. 3 | Projected Imitervirales" in sf)
check("legacy microdiversity Supplementary Fig. 3 is absent",
      "Supplementary Fig. 3 | Population microdiversity" not in si + sf)
check("internal historical figure notes are removed",
      "corrected in place" not in sf and "replaces Fig. S14" not in sf)

placeholder_count = len(re.findall(r"\[INSERT\]|\[GITHUB URL\]", ms))
check("only five known accession/DOI/URL placeholders remain", placeholder_count == 5, str(placeholder_count))
check("no AUTHOR VERIFY/TODO/FIXME markers", not re.search(r"AUTHOR VERIFY|TODO|FIXME", ms + si + gd + sf, re.I))

figs = sorted((ROOT / "figures").glob("*.pdf"))
check("five main and five Extended Data PDFs exist", len(figs) == 10, str(len(figs)))
archive = zipfile.ZipFile(MS)
embedded = {hashlib.md5(archive.read(n)).hexdigest() for n in archive.namelist() if "/media/" in n}
check("current figure PNGs are embedded in the manuscript",
      all(hashlib.md5((ROOT / "figures" / f"{p.stem}.png").read_bytes()).hexdigest() in embedded for p in figs))

failed = len([x for x in checks if not x])
print(f"\n{len(checks) - failed}/{len(checks)} checks passed")
sys.exit(1 if failed else 0)
