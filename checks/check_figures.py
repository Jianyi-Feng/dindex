"""Reproducible completeness and presentation checks for the  figures."""
import hashlib
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd
from PIL import Image
from docx import Document


ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
MS = ROOT / "manuscript/SPOT_giant_virus_Manuscript-.docx"
SF = ROOT / "supplementary/Supplementary_Figures.docx"

names = [f"Fig{i}_{x}" for i, x in enumerate(
    ["derive", "test", "validate", "interpret", "project"], 1)] + [
    f"ExtendedDataFig{i}_{x}" for i, x in enumerate(
        ["index_properties", "heatwave", "global_robustness", "projection_robustness", "marker_recovery"], 1)]
checks = []


def check(label, condition, detail=""):
    checks.append(bool(condition))
    print(f"[{'PASS' if condition else 'FAIL'}] {label}" + (f" — {detail}" if detail else ""))


missing = [f"{name}{ext}" for name in names for ext in (".png", ".pdf") if not (FIG / f"{name}{ext}").exists()]
check("all 10 Article/Extended Data figures have PNG and vector PDF", not missing, ", ".join(missing))

missing_source = [name for name in names if not (FIG / "source_data" / f"{name}_source_data.xlsx").exists()]
check("all 10 figures have machine-readable source workbooks", not missing_source, ", ".join(missing_source))

bad_sheets = []
for name in names:
    book = pd.ExcelFile(FIG / "source_data" / f"{name}_source_data.xlsx")
    if not book.sheet_names or any(pd.read_excel(book, sheet_name=s).empty for s in book.sheet_names):
        bad_sheets.append(name)
check("every source-data sheet is non-empty", not bad_sheets, ", ".join(bad_sheets))

sizes = {name: Image.open(FIG / f"{name}.png").size for name in names}
check("all Article/Extended Data PNGs are full-width 600-dpi exports",
      all(w >= 4200 and h >= 1100 for w, h in sizes.values()),
      "; ".join(f"{n}={w}×{h}" for n, (w, h) in sizes.items()))

with zipfile.ZipFile(MS) as z:
    media = {hashlib.md5(z.read(n)).hexdigest() for n in z.namelist() if "/media/" in n}
embedded = [name for name in names if hashlib.md5((FIG / f"{name}.png").read_bytes()).hexdigest() in media]
check("the manuscript embeds the current PNG for every figure", len(embedded) == 10, f"{len(embedded)}/10")

ms = Document(MS)
short_captions = [p.text for p in ms.paragraphs if re.match(r"^(?:Extended Data )?Fig\. \d+ \|", p.text)]
check("manuscript contains five main and five Extended Data figure pages",
      len([x for x in short_captions if x.startswith("Fig.")]) >= 5 and
      len([x for x in short_captions if x.startswith("Extended Data")]) >= 5)

sf = Document(SF)
sf_text = "\n".join(p.text for p in sf.paragraphs)
check("Supplementary figure document contains exactly Figs. 1–5",
      [int(x) for x in re.findall(r"Supplementary Fig\. (\d+) \|", sf_text)] == [1, 2, 3, 4, 5])
with zipfile.ZipFile(SF) as z:
    sf_media = [n for n in z.namelist() if "/media/" in n]
check("Supplementary document embeds five figure images", len(sf_media) == 5, str(len(sf_media)))

orientations = [(int(s.page_width) > int(s.page_height)) for s in sf.sections]
check("wide Supplementary figures use landscape pages and the network remains portrait",
      orientations == [False, True, True, True, True, False], str(orientations))

check("Imitervirales/Algavirales palette is declared consistently in the figure source",
      'I_C, A_C = "#0072B2", "#D55E00"' in (FIG / "make_figures.py").read_text())
check("Extended Data Fig. 4 uses environmental-support terminology",
      "environmental-support domain" in (FIG / "make_extended_data_fig4.py").read_text() and
      "area of applicability" not in (FIG / "make_extended_data_fig4.py").read_text())

failed = len([x for x in checks if not x])
print(f"\n{len(checks) - failed}/{len(checks)} checks passed")
sys.exit(1 if failed else 0)
