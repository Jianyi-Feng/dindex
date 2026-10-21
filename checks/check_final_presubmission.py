"""Final, submission-facing audit for the  Nature package.

Unlike check_package.py, this script treats unresolved author metadata and
promised-but-absent deliverables as blockers. It exits non-zero while any blocker
remains, which is intentional.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from scipy.stats import pearsonr


ROOT = Path(__file__).resolve().parent
MS = ROOT / "manuscript/SPOT_giant_virus_Manuscript-.docx"
CL = ROOT / "manuscript/Cover_Letter-Nature-.docx"
SI = ROOT / "supplementary/Supplementary_Information.docx"
GD = ROOT / "supplementary/SIGuide.docx"
SF = ROOT / "supplementary/Supplementary_Figures.docx"
SRC = ROOT / "processed_tables" / "figure_source_data"
results: list[tuple[str, str, str]] = []


def doc_text(path: Path) -> str:
    doc = Document(path)
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for table in doc.tables for row in table.rows for c in row.cells]
    )


def record(status: str, label: str, detail: str = "") -> None:
    results.append((status, label, detail))
    suffix = f" — {detail}" if detail else ""
    print(f"[{status}] {label}{suffix}")


def check(label: str, condition: bool, detail: str = "") -> None:
    record("PASS" if condition else "FAIL", label, detail)


def blocker(label: str, resolved: bool, detail: str = "") -> None:
    record("PASS" if resolved else "BLOCK", label, detail)


files = [MS, CL, SI, GD, SF]
check("five current  Word documents exist", all(p.exists() for p in files))
ms, cl, si, gd, sf = map(doc_text, files)

title = Document(MS).paragraphs[0].text
summary = ms[ms.index("Ocean warming is reorganizing"):ms.index("Marine species are shifting")]
summary_words = len(re.sub(r"\d", "", summary).split())
check("title meets Nature's 75-character guideline", len(title) <= 75, f"{len(title)} characters")
check("summary is at most 200 words", summary_words <= 200, f"{summary_words} extracted words")
check("main reference list is within the 50-reference guideline",
      len(re.findall(r"^\d+\. [A-Z]", ms[:ms.index("Figure legends")], re.M)) == 39,
      "39 main references; 71 total")

check("required disclosure/metadata sections are present",
      all(x in ms for x in ["Funding", "Use of large language models", "Additional information",
                            "Data availability", "Code availability"]))
check("funding is separated from acknowledgements",
      "This research was funded" not in ms[ms.index("Acknowledgements"):ms.index("Funding")])
check("LLM disclosure is inside the Methods sequence",
      ms.index("Use of large language models") < ms.index("Data availability") < ms.index("Methods references"))
check("correspondence statement is explicit",
      "Correspondence and requests for materials" in ms)

bad_markers = re.findall(r"AUTHOR VERIFY|TODO|FIXME|TBD|XXX", ms + cl + si + sf, re.I)
check("no internal drafting markers remain in scientific documents", not bad_markers,
      f"{len(bad_markers)} markers")

tables = sorted((ROOT / "supplementary").glob("Supplementary_Table_*.xlsx"))
table_nums = sorted({int(re.search(r"Table_(\d+)_", p.name).group(1)) for p in tables})
check("analysis-generated Supplementary Tables exist",
      table_nums == list(range(1, 11)) + list(range(12, 19)), str(table_nums))

workbooks = tables + sorted(SRC.glob("*source_data.xlsx"))
bad_workbooks = []
for path in workbooks:
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
        nonempty = sum(
            cell.value is not None
            for sheet in wb.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
        if not nonempty:
            bad_workbooks.append(f"{path.name}: empty")
    except Exception as exc:  # pragma: no cover - diagnostic path
        bad_workbooks.append(f"{path.name}: {exc}")
check("all current table/source-data workbooks open and contain data",
      not bad_workbooks, "; ".join(bad_workbooks) if bad_workbooks else f"{len(workbooks)} workbooks")

figure_pngs = sorted((ROOT / "figures").glob("Fig*.png")) + sorted((ROOT / "figures").glob("ExtendedDataFig*.png"))
figure_pdfs = sorted((ROOT / "figures").glob("Fig*.pdf")) + sorted((ROOT / "figures").glob("ExtendedDataFig*.pdf"))
check("five main and five Extended Data figures have PNG and vector PDF",
      len(figure_pngs) == 10 and len(figure_pdfs) == 10,
      f"{len(figure_pngs)} PNG; {len(figure_pdfs)} PDF")
check("all main/Extended Data figures have source workbooks",
      len([p for p in SRC.glob("*source_data.xlsx") if not p.name.startswith("Supplementary")]) == 10)

# Independent headline-statistic checks from the figure-source workbooks.
spot = pd.read_excel(SRC / "Fig1_derive_source_data.xlsx", sheet_name="a_samples")
share = spot["I"] + spot["A"]
check("SPOT sample counts and I+A dominance reproduce",
      len(spot) == 167 and spot["SST"].notna().sum() == 157
      and np.allclose(share.quantile([0.25, 0.5, 0.75]), [0.8837574, 0.96547833, 0.9900248], atol=1e-6))

tara = pd.read_excel(SRC / "Fig3_validate_source_data.xlsx", sheet_name="a_b_station_means")
r_temp, p_temp = pearsonr(tara["D"], tara["T"])
r_chl, _ = pearsonr(tara["D"], np.log10(tara["chl"]))
oof = pd.read_excel(SRC / "Fig3_validate_source_data.xlsx", sheet_name="c_oof_predictions")
r2 = 1 - ((oof.observed_D - oof.predicted_D) ** 2).sum() / ((oof.observed_D - oof.observed_D.mean()) ** 2).sum()
rmse = np.sqrt(np.mean((oof.observed_D - oof.predicted_D) ** 2))
check("Tara correlations and ensemble validation reproduce",
      len(tara) == 83 and len(oof) == 196 and np.isclose(r_temp, 0.7363065, atol=1e-6)
      and np.isclose(p_temp, 2.1596e-15, rtol=1e-3) and np.isclose(r_chl, -0.6495426, atol=1e-6)
      and np.isclose(r2, 0.6344520, atol=1e-6) and np.isclose(rmse, 0.4502231, atol=1e-6))

rates = pd.read_excel(SRC / "Fig5_project_source_data.xlsx", sheet_name="e_rates_by_projection")
primary = rates[rates["projection"].eq("final ensemble (7 predictors)")]
rate_means = primary.groupby("hemisphere")["km_per_decade"].mean()
maps = pd.read_excel(SRC / "Fig5_project_source_data.xlsx", sheet_name="a_b_c_maps_2deg_4model_mean")
maps = maps[maps["dD"].notna()].copy()
weights = np.cos(np.deg2rad(maps["lat"]))
positive_share = np.average(maps["dD"] > 0, weights=weights)
agreement_share = np.average(maps["all_models_agree"], weights=weights)
check("CMIP6 front rates and area-weighted map claims reproduce",
      np.isclose(rate_means["North"], 65.269323, atol=1e-6)
      and np.isclose(rate_means["South"], 36.027341, atol=1e-6)
      and np.isclose(positive_share, 0.972947, atol=1e-6)
      and np.isclose(agreement_share, 0.700406, atol=1e-6))

figure_audit = subprocess.run(
    [sys.executable, str(ROOT / "check_figures.py")],
    cwd=ROOT.parent.parent,
    capture_output=True,
    text=True,
)
check("dedicated figure audit passes", figure_audit.returncode == 0,
      "11/11" if figure_audit.returncode == 0 else figure_audit.stdout[-500:])

# Submission blockers: these are deliberately strict.
placeholder_count = len(re.findall(r"\[INSERT\]|\[GITHUB URL\]", ms))
blocker("all repository accessions, URLs and DOIs are final", placeholder_count == 0,
        f"{placeholder_count} unresolved fields")
reporting = list(ROOT.rglob("*Reporting*Summary*")) + list(ROOT.rglob("*reporting*summary*"))
blocker("completed Nature Reporting Summary is present", bool(reporting),
        "no Reporting Summary file found" if not reporting else reporting[0].name)
blocker("Supplementary Table 11 is complete",
        bool(list((ROOT / "supplementary").glob("Supplementary_Table_11_*.xlsx")))
        and "AUTHOR TO COMPLETE" not in gd,
        "software/database versions and critical parameters remain author-owned")
supp_data = [
    p for p in ROOT.rglob("Supplementary_Data_*")
    if p.is_file() and re.search(r"Supplementary_Data_(?:[1-9]|10)(?:\D|$)", p.name)
]
present_data_nums = sorted({
    int(re.search(r"Supplementary_Data_(\d+)", p.name).group(1)) for p in supp_data
})
blocker("Supplementary Data 1–10 inventoried in the SI are present",
        present_data_nums == list(range(1, 11)), f"found {present_data_nums}")

blockers = sum(status == "BLOCK" for status, _, _ in results)
failures = sum(status == "FAIL" for status, _, _ in results)
passes = sum(status == "PASS" for status, _, _ in results)
print(f"\n{passes} passed; {failures} failed; {blockers} submission blockers")
sys.exit(1 if blockers or failures else 0)
