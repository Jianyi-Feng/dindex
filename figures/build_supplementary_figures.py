"""Build the source-auditable Supplementary Figures (: five figures).

The per-fraction microdiversity figure carried over from the legacy package was dropped in : it
reversed the order colours, mislabelled the size fractions and reported sample-level statistics,
whereas Fig. 4c shows the per-genome comparison the Methods describe.
"""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from PIL import Image

HERE = Path(__file__).resolve().parent
PKG = HERE.parent
FIG = PKG / "outputs" / "supplementary"
OUT = PKG / "outputs" / "Supplementary_Figures.docx"
MAXW, MAXH = 165.0, 185.0

F = [
 ("SupplementaryFig1_fromS1.png", "Environmental and community context of the 2014–2016 Northeast Pacific marine heatwave at the San Pedro Ocean Time-series.",
  "a, Multivariate ENSO index (MEI), sea-surface temperature (SST), satellite chlorophyll a (Chla) and primary production (PP) at the SPOT location; the red box marks the heatwave. "
  "b, Relative abundance of giant-virus orders in the viral fraction (0.02–0.2 µm) in the pre-event period (2009–2013) and the early heatwave (2014–2015). "
  "c, The same order-level comparison in the micro/nanoplankton fraction (1.2–80 µm) between the heatwave (2014–2016) and subsequent years (2017–2019). "
  "d,e, Relative abundance of dominant eukaryotic 18S (d) and chloroplast 16S (e) classes before (2005–2013), during (2014–2016) and after (2017–2018) the heatwave. "
  "Differences were tested with Kruskal–Wallis tests followed by post-hoc pairwise comparisons with Benjamini–Hochberg adjustment (*P < 0.05, **P < 0.01, ***P < 0.001; ns, not significant). "
  "These panels provide order-level and host-community context only; the D-index event tests are reported in Fig. 2 and Supplementary Note 3. Note that the post-event window in c (2017–2019) is "
  "wider than the SST-matched post phase (2017–2018) used for the micro/nanoplankton D contrast in the main text. [AUTHOR VERIFY] Regenerate c on the 2017–2018 window, or state the difference in the "
  "final caption, so that the two windows cannot be confused. Supports Supplementary Notes 1 and 3."),
 ("SupplementaryFig2_fromS2.png", "Phytoplankton community composition at SPOT from chloroplast 16S rRNA gene sequences.",
  "Relative abundance of phytoplankton classes classified with PhytoREF (Methods) in the micro/nanoplankton fraction across the time series. These data provide the chloroplast 16S community "
  "matrix used in the Mantel analyses reported in Supplementary Table 2b and the community context in Supplementary Fig. 1e. Supports Supplementary Notes 1 and 4."),
 ("SupplementaryFig3_fromS3.png", "Species accumulation curves of giant-virus phylotypes.",
  "a, Accumulation curve for the SPOT dataset. b, Accumulation curve for the Tara Oceans stations. Both curves approach a plateau, indicating that the sampling effort captures the majority of "
  "detectable giant-virus PolB phylotype diversity in each dataset. Supports Supplementary Note 1."),
 ("SupplementaryFig4_fromS4.png", "Mantel associations between giant-virus community composition and environmental or biotic factors.",
  "Mantel r between giant-virus community dissimilarity in each size fraction (0.02–0.2 µm, 0.2–1.2 µm and 1.2–80 µm) and six environmental and biotic dissimilarity matrices. Circle area is "
  "proportional to Mantel r; orange-red, biotic factors; blue, environmental factors (*P < 0.05, **P < 0.01). Mantel tests are association measures and do not establish directionality, host "
  "identity or causal environmental control. Values are tabulated in Supplementary Table 2b. Supports Supplementary Note 4."),
 ("SupplementaryFig5_fromS7.png", "Phylogenomic placement of SPOT giant-virus and Mirusviricota genomes.",
  "a, Maximum-likelihood phylogeny of the Nucleocytoviricota inferred from a concatenated alignment of PolB, TFIIB, TopoII, A32, SFII, VLTF3 and RNAPL and rooted between Pokkesviricetes and "
  "Megaviricetes; the tree contains 1,620 tips (241 SPOT GVMAGs and 1,379 GVDB reference genomes). From the centre outwards, the first nine rings show the copy number of these markers plus MCP "
  "and RNAPS, followed by order-level taxonomy. Green triangles mark genomes recovered in this study; grey dots mark ultrafast bootstrap support > 95% (1,000 replicates); scale bar, one "
  "substitution per site. b, Maximum-likelihood phylogeny of Mirusviricota HK97-fold major capsid proteins, including the 25 SPOT bin-derived sequences; rings indicate clade and sequence source; "
  "purple dots mark bootstrap support > 95%. The phylogeny establishes order-level placement and shows that the compared Imitervirales and Algavirales genomes belong to deeply divergent clades; "
  "it is not used to claim independent evolutionary acquisition of the traits in Fig. 4. Three legend labels misspelt in the original figure file (“Pimscovirales”, “Alagvirales” and "
  "“Enodogenous”) were corrected in place; regenerate the panel from the editable source figure if it becomes available. Supports Supplementary Note 5.5 and Supplementary Table 11."),
 ("SupplementaryFig6_fromS8.png", "Annotated metabolic and stress-response genes across giant-virus orders and families.",
  "Heatmap of genes assigned to glycolysis, gluconeogenesis, the TCA cycle, the glyoxylate shunt, amino-acid metabolism, lipid metabolism, light harvesting, cytoskeleton, oxidative-stress "
  "regulation, DNA repair, ribosomal proteins, the ubiquitin system, apoptosis, cold-shock proteins and heat-shock proteins. Gene names and pathway assignments are listed in Supplementary "
  "Table 14. The five contrasts displayed in Fig. 4a are drawn from these categories; genome-level counts, Wilson intervals and Benjamini–Hochberg-adjusted tests for all 25 categories are given "
  "in Supplementary Table 8. Gene presence is an annotation-based contrast and does not establish expression or an effect on infection. Supports Supplementary Note 5."),
 ("SupplementaryFig7_fromS9.png", "Phylogenetic relationships of giant-virus rhodopsin sequences.",
  "Maximum-likelihood phylogeny rooted between type-I rhodopsins and heliorhodopsins. Sequences on the purple background are related to the reference green-light-absorbing proteorhodopsin (GPR) "
  "and blue-light-absorbing proteorhodopsin (BPR); sequences on the green background are predicted to have eight transmembrane helices. Scale bar, one substitution per site. In the final genome table, type-I rhodopsins were detected in 68 of 139 Imitervirales and 3 of 81 "
  "Algavirales genomes, and heliorhodopsin-like sequences in 27 of the 266 genomes analysed, predominantly Mirusviricota (18 genomes) and Pandoravirales "
  "(6 genomes; Supplementary Data 7). [AUTHOR VERIFY] The assignment of individual sequences to the GPR/BPR-like and eight-transmembrane groups shown "
  "here cannot be revalidated from the current data package; confirm it against the archived rhodopsin sequence catalogue. Sequence metadata are given in Supplementary Table 11. "
  "Supports Supplementary Note 5.6."),
 ("SupplementaryFig8_fromS10.png", "Structure prediction and sequence alignment of giant-virus rhodopsins.",
  "a, Amino-acid alignment of GPR, BPR and related sequences recovered here. b, Structural comparison of a Nucleocytoviricota rhodopsin with GPR. c, Predicted three-dimensional structure of a "
  "Mirusviricota heliorhodopsin-like sequence with eight transmembrane helices. d, Schematic of the eight-transmembrane heliorhodopsins identified here. Structures were predicted with AlphaFold 3 "
  "(Methods). These analyses give sequence-level context for the light-associated annotation category in Fig. 4a and are not required for the definition or calibration of D. "
  "Supports Supplementary Note 5.6."),
 ("SupplementaryFig9_fromS11.png", "Nucleotide-level selection signatures across giant-virus orders and pathways.",
  "a, Density of genome-wide pN/pS for Algavirales and Imitervirales populations in the micro/nanoplankton (1.2–80 µm) and viral (0.02–0.2 µm) fractions; the dashed line marks pN/pS = 1. "
  "b, pN/pS of selected pathway categories in the viral fraction, compared between orders with two-sided Mann–Whitney U tests and Benjamini–Hochberg adjustment (*P < 0.05). Values above 1 are "
  "compatible with relaxed purifying selection or with positive selection; this analysis cannot distinguish the two, and the main-text conclusions do not depend on it. [AUTHOR VERIFY] Insert final "
  "pathway-level sample sizes, coverage thresholds and adjusted P values from the archived inStrain output. Supports Supplementary Note 6.2."),
 ("SupplementaryFig10_fromS12.png", "Population microdiversity of the two orders by size fraction.",
  "Per-genome nucleotide diversity for Algavirales and Imitervirales populations in each size fraction, compared with two-sided Wilcoxon rank-sum tests (P values above each panel). The pooled "
  "order-level comparison is shown in Fig. 4c, and the per-genome summaries used to avoid pseudoreplication are given in Supplementary Table 10 and Supplementary Data 8. "
  "Supports Supplementary Note 6."),
 ("SupplementaryFig11_fromS13.png", "Co-occurrence associations between giant-virus phylotypes and eukaryotic ASVs.",
  "Association network inferred with FlashWeave from giant-virus PolB phylotypes and 18S eukaryotic ASVs present in at least 10% of micro/nanoplankton samples (P < 0.01, |weight| ≥ 0.4); edge "
  "thickness scales with the number of connections and red edges mark virus–host pairs known from cultured isolates. Edges are ecological associations that may reflect shared environmental "
  "responses, size fractionation or indirect interactions; they are not host assignments, and no main-text conclusion depends on them. [AUTHOR VERIFY] Insert final node and edge counts and "
  "deposit the node and edge tables as machine-readable Supplementary Data. Supports Supplementary Note 6.3."),
 ("SupplementaryFig12_projected_components.png", "Projected Imitervirales and Algavirales relative abundances underlying the D projections.",
  "a–c, Four-model mean predicted Imitervirales relative abundance in the 2000s (a) and 2090s under SSP5-8.5 (b), and the difference (c). d–f, The same fields for Algavirales. Predictions come "
  "from the final seven-predictor ensemble after removal of the mislabelled Tara pigment field, averaged over the picoplankton and femtoplankton calibrations and over four Earth-system models "
  "(CESM2-WACCM, IPSL-CM6A-LR, NorESM2-MM, UKESM1-0-LL); stippling in c and f marks cells where the four models disagree on the sign of the change. g, Zonal means by 1° latitude for both orders "
  "and both decades. h, Change per 15° latitude band. i, Decadal trajectories in the 45–60° N band for each Earth-system model. D combines these two fields; the equal-dominance fronts and the "
  "area of applicability of the calibration are shown in Fig. 5 and Extended Data Fig. 4. This figure replaces Fig. S14 of the earlier supplementary package, which was produced with the "
  "superseded eight-field projection. Supports Supplementary Note 7 and Supplementary Table 12."),
]

# Excludes legacy figures whose underlying matrices, sequence catalogues or result
# tables are absent from the current workspace, and (from ) the microdiversity
# figure superseded by Fig. 4c. Retained: phylogenomics, gene annotations, projections.
F = [F[i] for i in (4, 5, 11)]      # phylogenomics, gene heatmap, projected components
F[0] = (F[0][0], F[0][1], F[0][2].replace(
    " Three legend labels misspelt in the original figure file (“Pimscovirales”, “Alagvirales” and “Enodogenous”) were corrected in place; regenerate the panel from the editable source figure if it becomes available. Supports Supplementary Note 5.5 and Supplementary Table 11.",
    " Supports Supplementary Note 5.5; the Newick trees are included in the project archive."))
F[1] = (F[1][0], F[1][1], F[1][2].replace("Supplementary Table 14", "Supplementary Table 12"))
F[1] = (F[1][0], F[1][1], F[1][2].replace("Supplementary Table 8", "Supplementary Table 7"))
F[2] = (F[2][0], F[2][1], F[2][2].replace("Supplementary Table 12", "Supplementary Table 10"))
F[2] = (F[2][0], F[2][1], F[2][2].replace(
    " and the area of applicability of the calibration are shown in Fig. 5 and Extended Data Fig. 4. This figure replaces Fig. S14 of the earlier supplementary package, which was produced with the superseded eight-field projection.",
    " and the environmental-support domain of the calibration are shown in Fig. 5 and Extended Data Fig. 4."))
F.extend([
    ("SupplementaryFig4_pNpS.png", "pN/pS contrasts across analysis units, coverage filters and curated pathways.",
     "a, Median pN/pS of Imitervirales (blue) and Algavirales (vermilion) at genome level (solid; median of per-gene "
     "medians) and gene level (dashed; per-gene median across samples) under reporting tiers that filter gene-by-sample "
     "observations: A, coverage ≥10×; B, ≥5×; C, ≥1×; D, no filter; E, ≥10× with breadth ≥0.5. b, Cliff's δ "
     "(Imitervirales − Algavirales) with 95% CIs from bootstrap resampling of genomes; genomes (filled squares) are the "
     "primary unit because genes (open circles) are nested within genomes, so gene-level intervals use a genome-cluster "
     "bootstrap. The contrast is present at tiers A, B and E (genome δ = 0.37, 0.27 and 0.28) but not at tier C or D. "
     "c, Cliff's δ and 95% CIs for each curated pathway against the remaining curated genes in tier D (gene level, "
     "descriptive); filled black symbols mark Benjamini–Hochberg-adjusted P < 0.05. Values above one are compatible "
     "with relaxed purifying selection or positive selection and do not establish adaptation. Supports Supplementary "
     "Note 6.2 and Supplementary Table 14."),
    ("SupplementaryFig5_association_network.png", "Aggregated SPOT virus–eukaryote association network.",
     "The deposited network contains 31 nodes (six giant-virus orders and 25 eukaryotic 18S classes) and 53 "
     "order–class edges. Imitervirales are shown in blue, Algavirales in vermilion and the other giant-virus orders in "
     "grey; eukaryotic classes are teal. Edge width scales with the archived number of underlying associations (range "
     "1–41), and node size with weighted degree. The aggregate tables do not retain the original phylotype–ASV scores, "
     "signs or P values, so this figure is a descriptive summary rather than a re-estimation of the inference. "
     "Associations may reflect shared environmental responses, size fractionation or indirect interactions and are not "
     "host assignments. Supports Supplementary Note 6.3 and Supplementary Table 15."),
])

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.left_margin = s.right_margin = Mm(22); s.top_margin = s.bottom_margin = Mm(20)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15

h = doc.add_paragraph(); r = h.add_run("Supplementary Figures"); r.bold = True; r.font.size = Pt(13)
p = doc.add_paragraph()
p.add_run("A giant-virus index of warming-associated community reorganization. Supplementary Figures 1–5 accompany the Supplementary Information; Extended Data Figures 1–5 are submitted with the Article and are not repeated here.")
for i, (fn, title, body) in enumerate(F, 1):
    im = Image.open(FIG / fn)
    ratio = im.width / im.height
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    if ratio > 1.3:
        section.page_width, section.page_height = Mm(297), Mm(210)
        section.left_margin = section.right_margin = Mm(16)
        section.top_margin = section.bottom_margin = Mm(14)
        max_w, max_h = 260.0, 128.0
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(20)
        section.top_margin = section.bottom_margin = Mm(18)
        max_w, max_h = 170.0, 175.0
    w = min(max_w, max_h * ratio)
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(FIG / fn), width=Mm(w))
    cap = doc.add_paragraph()
    rr = cap.add_run(f"Supplementary Fig. {i} | {title}"); rr.bold = True
    cap.add_run(" " + body)
    cap.paragraph_format.space_after = Pt(12)

doc.core_properties.title = "Supplementary Figures"
doc.save(OUT)
print("wrote", OUT.name, "with", len(F), "figures")
